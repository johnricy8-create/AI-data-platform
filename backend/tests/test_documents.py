from io import BytesIO

import pandas as pd
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from app.services.documents import extract_document
from app.services.imports import _normalize_frame_columns


def _docx_with_table() -> bytes:
    document = Document()
    table = document.add_table(rows=2, cols=3)
    for index, value in enumerate(["订单号", "日期", "商品ID"]):
        table.rows[0].cells[index].text = value
    for index, value in enumerate(["DOCX-001", "2026-10-26", "P001"]):
        table.rows[1].cells[index].text = value
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pptx_with_table() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    shape = slide.shapes.add_table(2, 3, Inches(0.5), Inches(1.5), Inches(8), Inches(2))
    table = shape.table
    for index, value in enumerate(["订单号", "日期", "商品ID"]):
        table.cell(0, index).text = value
    for index, value in enumerate(["PPTX-001", "2026-10-26", "P001"]):
        table.cell(1, index).text = value
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def test_docx_table_is_extracted():
    result = extract_document("orders.docx", _docx_with_table())
    assert result.format == "docx"
    assert result.tables
    assert result.tables[0][1][0] == "DOCX-001"


def test_pptx_table_is_extracted():
    result = extract_document("orders.pptx", _pptx_with_table())
    assert result.format == "pptx"
    assert result.slides == 1
    assert result.tables
    assert result.tables[0][1][0] == "PPTX-001"


def test_chinese_excel_headers_are_normalized():
    frame = pd.DataFrame(
        [["XLSX-001", "2026-10-26", "淘宝", "已完成", "P001", 1, 399, 180]],
        columns=["订单号", "日期", "渠道", "状态", "商品ID", "数量", "单价", "成本价"],
    )
    normalized = _normalize_frame_columns(frame, "orders")
    assert {"order_id", "order_date", "channel", "status", "product_id", "quantity", "unit_price", "cost_price"} <= set(
        normalized.columns
    )
